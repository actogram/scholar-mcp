import json

from openpyxl import load_workbook

from actogram_scholar.export import to_csv, to_html, to_json, to_xlsx
from actogram_scholar.labs import build
from actogram_scholar.models import Author, Organization, Paper
from actogram_scholar.pipeline import Result

ORG = Organization(
    name="Uni <X> & Co",
    ror="https://ror.org/0x",
    city="Oslo",
    country="NO",
    country_name="Norway",
    lat=59.9,
    lon=10.7,
    match="ror",
)


def _result():
    papers = [
        Paper(
            title="Steps <script>alert(1)</script> and sleep",
            year=2025,
            pmid="123",
            authors=[Author("Nordmann", "Ola", affiliations=["x"])],
            mesh=["Humans"],
        )
    ]
    labs = build(papers, lambda aff: ORG, query="steps sleep")
    return Result(
        query="steps & sleep",
        start=2020,
        end=2025,
        sources=["pubmed"],
        papers=papers,
        labs=labs,
        counts={"pubmed": 1, "papers": 1, "labs": 1},
    )


def test_xlsx_has_labs_papers_and_about(tmp_path):
    wb = load_workbook(to_xlsx(_result(), tmp_path / "r.xlsx"))
    assert wb.sheetnames == ["Labs", "Papers", "About"]
    labs = wb["Labs"]
    assert labs["B2"].value == "Ola Nordmann" and labs["D2"].value == "Uni <X> & Co"
    assert wb["Papers"]["H2"].hyperlink.target == "https://pubmed.ncbi.nlm.nih.gov/123/"
    assert wb.properties.creator == "Actogram Scholar"


def test_html_escapes_text_and_keeps_the_data_readable(tmp_path):
    page = to_html(_result(), tmp_path / "r.html").read_text(encoding="utf-8")
    assert "<script>alert(1)</script>" not in page  # the title is inside JSON with "</" escaped
    assert "Labs for: steps &amp; sleep" in page
    data = json.loads(page.split('id="data">', 1)[1].split("</script>", 1)[0].replace("<\\/", "</"))
    assert data["labs"][0]["organization"] == "Uni <X> & Co" and data["labs"][0]["lat"] == 59.9
    assert 'integrity="sha256-' in page


def test_json_and_csv(tmp_path):
    data = json.loads(to_json(_result(), tmp_path / "r.json").read_text(encoding="utf-8"))
    assert data["query"] == "steps & sleep" and data["labs"][0]["lead"] == "Ola Nordmann"
    lines = to_csv(_result(), tmp_path / "r.csv").read_text(encoding="utf-8").splitlines()
    assert lines[0].startswith("rank,lead,orcid,organization") and "Ola Nordmann" in lines[1]
