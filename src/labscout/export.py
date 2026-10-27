"""Results as JSON, CSV, Excel, and a single HTML file with a map and a table that can be sorted and filtered."""

from __future__ import annotations

import csv
import html
import json
from collections import Counter
from pathlib import Path

from .pipeline import LABELS, Result, read_as

SOURCES_NOTE = (
    "Sources: PubMed (US National Library of Medicine, through NCBI E-utilities), Europe PMC, Research "
    "Organization Registry (ROR, CC0), NIH RePORTER and ClinicalTrials.gov. Names and affiliations are shown "
    "as published; e-mail addresses are removed. Organisations are matched automatically and can be wrong: "
    "the column 'match' says how each was found."
)
LAB_COLUMNS = [
    "rank",
    "lead",
    "orcid",
    "organization",
    "city",
    "country",
    "ror",
    "match",
    "papers_as_senior_author",
    "papers_as_coauthor",
    "first_year",
    "last_year",
    "species",
    "data_signals",
    "matched_terms",
    "funders",
    "nih_grants",
    "why",
    "top_paper",
    "top_paper_url",
]
PAPER_COLUMNS = [
    "year",
    "title",
    "journal",
    "first_author",
    "senior_author",
    "pmid",
    "doi",
    "url",
    "preprint",
    "species_mesh",
    "sources",
]


def lab_rows(result: Result) -> list[dict[str, object]]:
    rows = []
    for i, lab in enumerate(result.labs, 1):
        org, top = lab.organization, (lab.papers[0] if lab.papers else None)
        rows.append(
            {
                "rank": i,
                "lead": lab.lead,
                "orcid": lab.orcid,
                "organization": org.name,
                "city": org.city,
                "country": org.country_name or org.country,
                "ror": org.ror,
                "match": org.match,
                "papers_as_senior_author": lab.senior_papers,
                "papers_as_coauthor": lab.other_papers,
                "first_year": lab.first_year,
                "last_year": lab.last_year,
                "species": ", ".join(lab.species),
                "data_signals": ", ".join(lab.signals),
                "matched_terms": ", ".join(lab.matched_terms),
                "funders": ", ".join(lab.funders[:5]),
                "nih_grants": "; ".join(f"{g.number} {g.title}" for g in lab.grants),
                "why": lab.why,
                "top_paper": top.title if top else "",
                "top_paper_url": top.url if top else "",
            }
        )
    return rows


def paper_rows(result: Result) -> list[dict[str, object]]:
    rows = []
    for p in sorted(result.papers, key=lambda x: -(x.year or 0)):
        rows.append(
            {
                "year": p.year,
                "title": p.title,
                "journal": p.journal,
                "first_author": p.authors[0].display if p.authors else "",
                "senior_author": p.authors[-1].display if p.authors else "",
                "pmid": p.pmid,
                "doi": p.doi,
                "url": p.url,
                "preprint": "yes" if p.is_preprint else "",
                "species_mesh": ", ".join(
                    m
                    for m in p.mesh
                    if m in ("Humans", "Mice", "Rats", "Animals", "Drosophila melanogaster", "Zebrafish")
                ),
                "sources": ", ".join(p.sources),
            }
        )
    return rows


def to_json(result: Result, path: str | Path, top: int | None = None) -> Path:
    path = Path(path)
    path.write_text(
        json.dumps(result.summary(top=top or len(result.labs)), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return path


def to_csv(result: Result, path: str | Path) -> Path:
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LAB_COLUMNS)
        w.writeheader()
        w.writerows(lab_rows(result))
    return path


def to_xlsx(result: Result, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    path = Path(path)
    wb = Workbook()
    wb.properties.creator = wb.properties.lastModifiedBy = "Actogram LabScout"
    wb.properties.title = f"Labs for: {result.query}"
    for name, columns, rows in (("Labs", LAB_COLUMNS, lab_rows(result)), ("Papers", PAPER_COLUMNS, paper_rows(result))):
        ws = wb.active if name == "Labs" else wb.create_sheet(name)
        ws.title = name
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        link_col = columns.index("top_paper_url" if name == "Labs" else "url") + 1
        for row in ws.iter_rows(min_row=2, min_col=link_col, max_col=link_col):
            for cell in row:
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = Font(color="006956", underline="single")
        for i, c in enumerate(columns, 1):
            wide = {"why": 70, "title": 70, "top_paper": 60, "organization": 40, "lead": 24}.get(c, 14)
            ws.column_dimensions[get_column_letter(i)].width = wide
            if wide >= 40:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    about = wb.create_sheet("About")
    for line in (
        ["Search", result.query],
        ["Years", years_text(result)],
        ["Sources searched", ", ".join(result.sources)],
        ["Papers read", "the most relevant" if result.order == "relevance" else "the newest"],
        ["Found", ", ".join(f"{k}: {v}" for k, v in result.found.items())],
        ["Counts", ", ".join(f"{k}: {v}" for k, v in result.counts.items())],
        ["PubMed read the search as", result.translation],
        ["Created", result.created],
        ["Tool", f"Actogram LabScout {result.version}, https://github.com/actogram/scholar-mcp"],
        ["Notes", SOURCES_NOTE],
    ):
        about.append(line)
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 110
    for row in about.iter_rows(min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return path


def _n(count: int, word: str) -> str:
    return f"{count:,} {word}" + ("" if count == 1 else "s")


def years_text(result: Result) -> str:
    s, e = result.start, result.end
    if s and e:
        return f"{s} to {e}"
    if s:
        return f"{s} to now"
    return f"Up to {e}" if e else "All years"


def overview(result: Result) -> list[str]:
    """Plain sentences for the top of the report: what the sources have, what was read and found, and what is
    missing."""
    out = []
    if not result.papers:
        names = " and ".join(LABELS.get(s, s) for s in result.sources if s not in result.failed)
        return [f"{years_text(result)}: {names or 'The sources'} found no papers for this search."]
    have = [(LABELS.get(s, s), result.found[s]) for s in result.sources if s in result.found]
    if have:
        said = " and ".join(f"{name} has {n:,}" for name, n in have)
        out.append(f"{years_text(result)}: {said} {'paper' if have[-1][1] == 1 else 'papers'} for this search.")
    read = [
        f"the {result.counts.get(s, 0)} {read_as(s, result.order)} from {LABELS.get(s, s)}"
        for s in result.sources
        if s not in result.failed
    ]
    papers = result.counts.get("papers", len(result.papers))
    if read:
        out.append(
            f"LabScout read {' and '.join(read)}" + (f": {_n(papers, 'different paper')}." if len(read) > 1 else ".")
        )
    countries = len({lab.organization.country_name or lab.organization.country for lab in result.labs} - {""})
    out.append(f"It found {_n(len(result.labs), 'lab')} in {countries} {'country' if countries == 1 else 'countries'}.")
    if result.failed:
        names = " and ".join(LABELS.get(s, s) for s in result.failed)
        others = " and ".join(LABELS.get(s, s) for s in result.sources if s not in result.failed)
        out.append(
            f"{names} did not answer, so these results come from {others or 'no source'} only. Please try again later."
        )
    unplaced = sum(1 for lab in result.labs if lab.organization.lat is None)
    if unplaced and result.labs:
        out.append(f"{unplaced} of the labs have no known place: they are in the table but not on the map.")
    return out


def hint(result: Result) -> str:
    """Advice when the search is too broad or too narrow, or empty."""
    papers = result.counts.get("papers", len(result.papers))
    if papers == 0:
        return "No papers match this search. Use fewer words, OR between similar words, or more years."
    if papers < 20:
        return f"Only {_n(papers, 'paper')} match. For more, use fewer words, OR between similar words, or more years."
    broad = any(result.found.get(s, 0) > 5 * max(1, result.counts.get(s, 0)) for s in result.sources)
    if broad:
        which = "newest" if result.order == "newest" else "most relevant and most cited"
        return (
            f"This topic is broad, so the labs come from the {which} papers only. To focus, add a word, put a phrase "
            "in quotes or narrow the years."
        )
    return ""


def country_bars(result: Result, top: int = 8) -> str:
    """Labs by country: one bar each for the countries with the most labs, the rest together as Other."""
    counts = Counter(lab.organization.country_name or lab.organization.country for lab in result.labs)
    counts.pop("", None)
    if not counts:
        return ""
    rows = counts.most_common(top)
    rest = sum(counts.values()) - sum(n for _, n in rows)
    if rest:
        rows.append(("Other", rest))
    most = max(n for _, n in rows)
    items = "".join(
        f'<li title="{html.escape(name)}: {n} lab{"s" if n != 1 else ""}"><span class="k">{html.escape(name)}</span>'
        f'<span class="b"><i style="width:{max(2.0, 100 * n / most):.1f}%"></i></span><span class="v">{n}</span></li>'
        for name, n in rows
    )
    return f'<figure class="chart"><figcaption>Labs by country</figcaption><ol class="hbars">{items}</ol></figure>'


def year_columns(result: Result, bars: int = 15) -> str:
    """Papers by year, as small columns: the most recent years, the rest together as Earlier."""
    counts = Counter(p.year for p in result.papers if p.year)
    if not counts:
        return ""
    years = sorted(counts)
    shown = years[-bars:]
    cols = [(str(y), counts[y]) for y in shown]
    earlier = sum(counts[y] for y in years[: -bars or None]) if len(years) > bars else 0
    if earlier:
        cols.insert(0, ("Earlier", earlier))
    w, h, top, base = 360, 150, 18, 124
    band = w / len(cols)
    bar = min(20.0, band * 0.6)
    most = max(n for _, n in cols)
    label_all = len(cols) <= 8
    peak = max(range(len(cols)), key=lambda i: cols[i][1])
    parts = [f'<line x1="0" y1="{base}" x2="{w}" y2="{base}" class="base"/>']
    for i, (name, n) in enumerate(cols):
        x = i * band + (band - bar) / 2
        hgt = max(2.0, (base - top) * n / most)
        y = base - hgt
        r = min(4.0, bar / 2, hgt)
        d = f"M{x:.1f} {base}V{y + r:.1f}Q{x:.1f} {y:.1f} {x + r:.1f} {y:.1f}H{x + bar - r:.1f}"
        d += f"Q{x + bar:.1f} {y:.1f} {x + bar:.1f} {y + r:.1f}V{base}Z"
        cx = x + bar / 2
        label = html.escape(name)
        parts.append(f'<g><title>{label}: {n} paper{"s" if n != 1 else ""}</title><path d="{d}"/>')
        if label_all or i in (peak, len(cols) - 1):
            parts.append(f'<text x="{cx:.1f}" y="{y - 5:.1f}" class="v">{n}</text>')
        if label_all or i % 2 == (len(cols) - 1) % 2:
            short = name if label_all or name == "Earlier" else "\u2019" + name[2:]
            parts.append(f'<text x="{cx:.1f}" y="{base + 17}" class="k">{html.escape(short)}</text>')
        parts.append("</g>")
    svg = (
        f'<svg viewBox="0 0 {w} {h}" class="cols" role="img" aria-label="Papers by year, '
        + html.escape(", ".join(f"{a}: {b}" for a, b in cols))
        + '">'
        + "".join(parts)
        + "</svg>"
    )
    return f'<figure class="chart"><figcaption>Papers by year</figcaption>{svg}</figure>'


def render_html(result: Result, top: int = 200, toolbar: str = "", footer: str = "") -> str:
    """The report page as text; `toolbar` is HTML placed under the title and `footer` at the end (the web page puts its
    links there)."""
    data = result.summary(top=top, papers_per_lab=5)
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    e = html.escape
    advice = hint(result)
    reading = (
        f'<details class="reading"><summary>How PubMed read your search</summary><code>{e(result.translation)}</code>'
        "</details>"
        if result.translation
        else ""
    )
    charts = country_bars(result) + year_columns(result)
    values = {
        "{{TITLE}}": e(f"Labs for: {result.query}"),
        "{{DESC}}": e(f"Labs working on {result.query}, on a map and in a table, from Actogram LabScout."),
        "{{QUERY}}": e(result.query),
        "{{OVERVIEW}}": "".join(f"<p>{e(line)}</p>" for line in overview(result)),
        "{{HINT}}": f'<p class="hint">{e(advice)}</p>' if advice else "",
        "{{READING}}": reading,
        "{{CHARTS}}": f'<div class="charts">{charts}</div>' if charts else "",
        "{{CREATED}}": e(result.created),
        "{{NOTE}}": e(SOURCES_NOTE),
        "{{VERSION}}": e(result.version),
        "{{TOOLBAR}}": toolbar,
        "{{FOOTER}}": footer,
    }
    page = TEMPLATE
    for k, v in values.items():
        page = page.replace(k, v)
    return page.replace("{{DATA}}", payload)


def to_html(result: Result, path: str | Path, top: int = 200) -> Path:
    path = Path(path)
    path.write_text(render_html(result, top), encoding="utf-8")
    return path


def xlsx_bytes(result: Result) -> bytes:
    import io
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        out = to_xlsx(result, Path(tmp) / "labs.xlsx")
        return io.BytesIO(out.read_bytes()).getvalue()


def csv_text(result: Result) -> str:
    import io

    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=LAB_COLUMNS)
    w.writeheader()
    w.writerows(lab_rows(result))
    return buf.getvalue()


TEMPLATE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{TITLE}}</title>
<meta name="description" content="{{DESC}}">
<link rel="icon" href="data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 32 32'%3E%3Ccircle cx='13' cy='13' r='8.5' fill='none' stroke='%23006956' stroke-width='3.5'/%3E%3Cpath d='M19.5 19.5l7 7' stroke='%23006956' stroke-width='4' stroke-linecap='round'/%3E%3C/svg%3E">
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css"
  integrity="sha256-p4NxAoJBhIIN+hmNHrzRCf9tD/miZyoHS5obTRR9BMY=" crossorigin="">
<style>
:root { --ink:#1d2b27; --soft:#4f5a56; --green:#006956; --deep:#00473a; --mint:#d8f8f2; --tint:#f1f8f5; --line:#dfe8e4; --bg:#fff; }
* { box-sizing:border-box; }
body { margin:0; font:16px/1.5 system-ui, -apple-system, "Segoe UI", Roboto, sans-serif; color:var(--ink); background:var(--bg); }
header, main, footer { max-width:1200px; margin:0 auto; padding:0 20px; }
header { padding-top:28px; }
h1 { font-size:28px; line-height:1.2; margin:0 0 10px; }
.overview { max-width:880px; color:var(--soft); font-size:15.5px; }
.overview p { margin:0 0 6px; }
.hint { max-width:880px; margin:10px 0; padding:10px 14px; border:1px solid var(--line); border-radius:10px; background:var(--tint); font-size:15px; }
.reading { max-width:880px; margin:8px 0 4px; color:var(--soft); font-size:14.5px; }
.reading summary { cursor:pointer; color:var(--green); }
.reading code { display:block; margin-top:6px; padding:8px 10px; border-radius:8px; background:var(--tint);
  font:13px/1.5 ui-monospace, SFMono-Regular, Menlo, Consolas, monospace; overflow-wrap:anywhere; }
.tools { display:flex; flex-wrap:wrap; gap:8px; margin:14px 0 18px; font-size:15px; }
.tools a { display:inline-block; padding:5px 12px; border:1px solid var(--line); border-radius:8px; text-decoration:none; }
.tools a:hover { background:var(--tint); }
.charts { display:grid; grid-template-columns:1fr 1fr; gap:16px; margin:0 0 18px; }
.chart { margin:0; padding:14px 16px; border:1px solid var(--line); border-radius:12px; }
.chart figcaption { margin:0 0 10px; font-weight:600; font-size:15px; }
.hbars { list-style:none; margin:0; padding:0; display:grid; gap:5px; }
.hbars li { display:grid; grid-template-columns:minmax(6em, 11em) 1fr 2.6em; gap:8px; align-items:center; font-size:14px; }
.hbars .k { overflow-wrap:anywhere; }
.hbars .b i { display:block; height:12px; background:var(--green); border-radius:0 4px 4px 0; }
.hbars li:hover .b i { background:var(--deep); }
.hbars .v { text-align:right; color:var(--soft); font-variant-numeric:tabular-nums; }
.cols { display:block; width:100%; height:auto; }
.cols path { fill:var(--green); }
.cols g:hover path { fill:var(--deep); }
.cols .base { stroke:var(--line); stroke-width:1; }
.cols text { font:11px system-ui, -apple-system, "Segoe UI", Roboto, sans-serif; fill:var(--soft); text-anchor:middle; }
#map { height:420px; margin:0 0 18px; border:1px solid var(--line); border-radius:12px; background:var(--tint); }
.bar { display:flex; flex-wrap:wrap; gap:10px; align-items:center; margin:0 0 10px; }
.bar input { flex:1 1 260px; padding:9px 12px; font:inherit; border:1px solid var(--line); border-radius:8px; }
.bar span { color:var(--soft); font-size:14px; }
.wrap { overflow-x:auto; border:1px solid var(--line); border-radius:12px; }
table { width:100%; border-collapse:collapse; font-size:14.5px; }
th, td { text-align:left; vertical-align:top; padding:9px 10px; border-bottom:1px solid var(--line); }
th { position:sticky; top:0; background:var(--mint); cursor:pointer; white-space:nowrap; }
th[aria-sort="ascending"]::after { content:" \25B2"; } th[aria-sort="descending"]::after { content:" \25BC"; }
td.num { text-align:right; white-space:nowrap; }
details summary { cursor:pointer; color:var(--green); }
details ol { margin:6px 0 0 18px; padding:0; }
a { color:var(--green); }
.why { color:var(--soft); }
footer { padding-top:18px; padding-bottom:28px; color:var(--soft); font-size:13.5px; }
@media (max-width:760px) { .charts { grid-template-columns:1fr; } }
@media (max-width:640px) {
  #map { height:300px; } h1 { font-size:23px; }
  .wrap { overflow:visible; border:0; }
  #labs thead { display:none; }
  #labs, #labs tbody, #labs tr, #labs td { display:block; width:100%; }
  #labs tr { margin:0 0 10px; padding:8px 2px; border:1px solid var(--line); border-radius:12px; }
  #labs td { padding:3px 12px; border:0; }
  #labs td.num { text-align:left; }
  #labs td.rank { font-weight:600; }
  #labs td.rank::before { content:"#"; }
  #labs td[data-label]::before { content:attr(data-label) ": "; color:var(--soft); font-size:13.5px; }
}
</style>
</head>
<body>
<header>
  <h1>Labs for: {{QUERY}}</h1>
  <div class="overview">{{OVERVIEW}}</div>
  {{HINT}}
  {{READING}}
  {{TOOLBAR}}
</header>
<main>
  {{CHARTS}}
  <div id="map" role="img" aria-label="Map of the labs that have a known place"></div>
  <div class="bar"><input id="q" type="search" placeholder="Filter by name, place, species or data" aria-label="Filter the labs">
    <span id="shown" aria-live="polite"></span></div>
  <div class="wrap"><table id="labs">
    <thead><tr><th data-k="rank">#</th><th data-k="lead">Lead</th><th data-k="organization">Organisation</th>
      <th data-k="country">Country</th><th data-k="papers_as_senior_author">Papers led</th><th data-k="last">Last paper</th>
      <th data-k="why">Why it came up</th></tr></thead>
    <tbody></tbody>
  </table></div>
</main>
<footer>{{FOOTER}}<p>Created {{CREATED}}. {{NOTE}}</p><p>Map tiles: &copy; OpenStreetMap contributors. Made with Actogram LabScout
  {{VERSION}}, <a href="https://github.com/actogram/scholar-mcp">github.com/actogram/scholar-mcp</a>.</p></footer>
<script type="application/json" id="data">{{DATA}}</script>
<script>
(function () {
  var data = JSON.parse(document.getElementById('data').textContent);
  var labs = data.labs.map(function (l, i) { l.rank = i + 1; l.last = (l.years || [])[1] || ''; return l; });
  function esc(s) { return String(s == null ? '' : s).replace(/[&<>"']/g, function (c) {
    return {'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'}[c]; }); }
  function papers(l) {
    if (!l.top_papers.length) return '';
    var n = l.top_papers.length;
    return '<details><summary>' + n + (n === 1 ? ' paper' : ' papers') + '</summary><ol>' + l.top_papers.map(function (p) {
      return '<li>' + (p.url ? '<a href="' + esc(p.url) + '" target="_blank" rel="noopener">' + esc(p.title) + '</a>' : esc(p.title)) +
        ' (' + esc(p.year) + ')</li>'; }).join('') + '</ol></details>';
  }
  var body = document.querySelector('#labs tbody');
  function row(l) {
    var lead = esc(l.lead) + (l.orcid ? ' <a href="https://orcid.org/' + esc(l.orcid) + '" target="_blank" rel="noopener">ORCID</a>' : '');
    var org = l.ror ? '<a href="' + esc(l.ror) + '" target="_blank" rel="noopener">' + esc(l.organization) + '</a>' : esc(l.organization);
    return '<tr><td class="num rank">' + l.rank + '</td><td data-label="Lead">' + lead + '</td><td data-label="Organisation">' + org +
      (l.city ? '<br><span class="why">' + esc(l.city) + '</span>' : '') + '</td><td data-label="Country">' + esc(l.country) +
      '</td><td class="num" data-label="Papers led">' + l.papers_as_senior_author + '</td><td class="num" data-label="Last paper">' +
      esc(l.last) + '</td><td><span class="why">' + esc(l.why) + '</span>' + papers(l) + '</td></tr>';
  }
  var key = 'rank', dir = 1;
  function draw() {
    var q = document.getElementById('q').value.toLowerCase();
    var list = labs.filter(function (l) { return !q || [l.lead, l.organization, l.city, l.country, l.why].join(' ').toLowerCase().indexOf(q) >= 0; });
    list.sort(function (a, b) { var x = a[key], y = b[key]; return (x > y ? 1 : x < y ? -1 : 0) * dir; });
    body.innerHTML = list.map(row).join('');
    document.getElementById('shown').textContent = list.length + ' of ' + labs.length + ' labs';
  }
  document.querySelectorAll('#labs th').forEach(function (th) {
    th.addEventListener('click', function () {
      var k = th.getAttribute('data-k'); dir = (k === key) ? -dir : (k === 'rank' || k === 'lead' ? 1 : -1); key = k;
      document.querySelectorAll('#labs th').forEach(function (h) { h.removeAttribute('aria-sort'); });
      th.setAttribute('aria-sort', dir > 0 ? 'ascending' : 'descending'); draw();
    });
  });
  document.getElementById('q').addEventListener('input', draw);
  draw();
  window.labscoutMap = function () {
    if (!window.L) {
      document.getElementById('map').textContent = 'The map needs an internet connection; the table works without it.';
      return;
    }
    var map = L.map('map', { scrollWheelZoom: false }).setView([30, 10], 2);
    L.tileLayer('https://tile.openstreetmap.org/{z}/{x}/{y}.png', { maxZoom: 18,
      attribution: '&copy; <a href="https://www.openstreetmap.org/copyright">OpenStreetMap</a> contributors' }).addTo(map);
    var pts = [];
    labs.forEach(function (l) {
      if (l.lat == null || l.lon == null) return;
      var r = 5 + Math.min(10, 2 * l.papers_as_senior_author);
      L.circleMarker([l.lat, l.lon], { radius: r, color: '#00473a', weight: 1.5, fillColor: '#006956', fillOpacity: 0.55 })
        .bindPopup('<b>' + esc(l.lead) + '</b><br>' + esc(l.organization) + '<br>' + esc(l.why)).addTo(map);
      pts.push([l.lat, l.lon]);
    });
    if (pts.length > 1) map.fitBounds(pts, { padding: [24, 24], maxZoom: 6 });
    else if (pts.length === 1) map.setView(pts[0], 5);
  };
})();
</script>
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js" integrity="sha256-20nQCchB9co0qIjJZRGuk2/Z9VM+kNiyxNV1lvTlZBo="
  crossorigin="" defer></script>
<script>document.addEventListener('DOMContentLoaded', function () { window.labscoutMap(); });</script>
</body>
</html>
"""
