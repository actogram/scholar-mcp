"""Results as JSON, CSV, Excel, and a single HTML file with a map and a table that can be sorted and filtered."""

from __future__ import annotations

import csv
import html
import json
from pathlib import Path

from .pipeline import Result

SOURCES_NOTE = (
    "Sources: PubMed (US National Library of Medicine, through NCBI E-utilities), Europe PMC, Research "
    "Organization Registry (ROR, CC0), NIH RePORTER and ClinicalTrials.gov. Names and affiliations are shown "
    "as published; e-mail addresses are removed. Organisations are matched automatically and can be wrong: "
    "the column 'match' says how each was found."
)
LAB_COLUMNS = [
    "rank",
    "lead",
    "orcid",
    "organization",
    "city",
    "country",
    "ror",
    "match",
    "papers_as_senior_author",
    "papers_as_coauthor",
    "first_year",
    "last_year",
    "species",
    "data_signals",
    "matched_terms",
    "funders",
    "nih_grants",
    "why",
    "top_paper",
    "top_paper_url",
]
PAPER_COLUMNS = [
    "year",
    "title",
    "journal",
    "first_author",
    "senior_author",
    "pmid",
    "doi",
    "url",
    "preprint",
    "species_mesh",
    "sources",
]


def lab_rows(result: Result) -> list[dict[str, object]]:
    rows = []
    for i, lab in enumerate(result.labs, 1):
        org, top = lab.organization, (lab.papers[0] if lab.papers else None)
        rows.append(
            {
                "rank": i,
                "lead": lab.lead,
                "orcid": lab.orcid,
                "organization": org.name,
                "city": org.city,
                "country": org.country_name or org.country,
                "ror": org.ror,
                "match": org.match,
                "papers_as_senior_author": lab.senior_papers,
                "papers_as_coauthor": lab.other_papers,
                "first_year": lab.first_year,
                "last_year": lab.last_year,
                "species": ", ".join(lab.species),
                "data_signals": ", ".join(lab.signals),
                "matched_terms": ", ".join(lab.matched_terms),
                "funders": ", ".join(lab.funders[:5]),
                "nih_grants": "; ".join(f"{g.number} {g.title}" for g in lab.grants),
                "why": lab.why,
                "top_paper": top.title if top else "",
                "top_paper_url": top.url if top else "",
            }
        )
    return rows


def paper_rows(result: Result) -> list[dict[str, object]]:
    rows = []
    for p in sorted(result.papers, key=lambda x: -(x.year or 0)):
        rows.append(
            {
                "year": p.year,
                "title": p.title,
                "journal": p.journal,
                "first_author": p.authors[0].display if p.authors else "",
                "senior_author": p.authors[-1].display if p.authors else "",
                "pmid": p.pmid,
                "doi": p.doi,
                "url": p.url,
                "preprint": "yes" if p.is_preprint else "",
                "species_mesh": ", ".join(
                    m
                    for m in p.mesh
                    if m in ("Humans", "Mice", "Rats", "Animals", "Drosophila melanogaster", "Zebrafish")
                ),
                "sources": ", ".join(p.sources),
            }
        )
    return rows


def to_json(result: Result, path: str | Path, top: int | None = None) -> Path:
    path = Path(path)
    path.write_text(
        json.dumps(result.summary(top=top or len(result.labs)), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return path


def to_csv(result: Result, path: str | Path) -> Path:
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LAB_COLUMNS)
        w.writeheader()
        w.writerows(lab_rows(result))
    return path


def to_xlsx(result: Result, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    path = Path(path)
    wb = Workbook()
    wb.properties.creator = wb.properties.lastModifiedBy = "Actogram Scholar"
    wb.properties.title = f"Labs for: {result.query}"
    for name, columns, rows in (("Labs", LAB_COLUMNS, lab_rows(result)), ("Papers", PAPER_COLUMNS, paper_rows(result))):
        ws = wb.active if name == "Labs" else wb.create_sheet(name)
        ws.title = name
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        link_col = columns.index("top_paper_url" if name == "Labs" else "url") + 1
        for row in ws.iter_rows(min_row=2, min_col=link_col, max_col=link_col):
            for cell in row:
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = Font(color="006956", underline="single")
        for i, c in enumerate(columns, 1):
            wide = {"why": 70, "title": 70, "top_paper": 60, "organization": 40, "lead": 24}.get(c, 14)
            ws.column_dimensions[get_column_letter(i)].width = wide
            if wide >= 40:
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
    about = wb.create_sheet("About")
    for line in (
        ["Search", result.query],
        ["Years", f"{result.start or ''} to {result.end or ''}"],
        ["Sources searched", ", ".join(result.sources)],
        ["Counts", ", ".join(f"{k}: {v}" for k, v in result.counts.items())],
        ["Created", result.created],
        ["Tool", f"Actogram Scholar {result.version}, https://github.com/actogram/scholar-mcp"],
        ["Notes", SOURCES_NOTE],
    ):
        about.append(line)
    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 110
    for row in about.iter_rows(min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return path


def to_html(result: Result, path: str | Path, top: int = 200) -> Path:
    path = Path(path)
    data = result.summary(top=top, papers_per_lab=5)
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    e = html.escape
    years = " to ".join(str(y) for y in (result.start, result.end) if y) or "all years"
    page = (
        TEMPLATE.replace("{{TITLE}}", e(f"Labs for: {result.query}"))
        .replace("{{QUERY}}", e(result.query))
        .replace("{{YEARS}}", e(years))
        .replace("{{COUNTS}}", e(", ".join(f"{k}: {v}" for k, v in result.counts.items())))
        .replace("{{CREATED}}", e(result.created))
        .replace("{{NOTE}}", e(SOURCES_NOTE))
        .replace("{{VERSION}}", e(result.version))
        .replace("{{DATA}}", payload)
    )
    path.write_text(page, encoding="utf-8")
    return path


TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{TITLE}}</title>
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css"
  integrity="sha256-p4NxAoJBhIIN+hmNHrzRCf9tD/miZyoHS5obTRR9BMY=" crossorigin="">
<style>
:root { --ink:#1d2b27; --soft:#4f5a56; --green:#006956; --mint:#d8f8f2; --line:#dfe8e4; --bg:#fff; }
* { box-sizing:border-box; }
body { margin:0; font:16px/1.5 system-ui, -apple-system, "Segoe UI", Roboto, sans-serif; color:var(--ink); background:var(--bg); }
header, main, footer { max-width:1200px; margin:0 auto; padding:0 20px; }
header { padding-top:28px; }
h1 { font-size:28px; line-height:1.2; margin:0 0 6px; }
.meta { color:var(--soft); font-size:15px; margin:0 0 18px; }
#map { height:420px; border:1px solid var(--line); border-radius:12px; margin:0 0 18px; }
.bar { display:flex; flex-wrap:wrap; gap:10px; align-items:center; margin:0 0 10px; }
.bar input { flex:1 1 260px; padding:9px 12px; font:inherit; border:1px solid var(--line); border-radius:8px; }
.bar span { color:var(--soft); font-size:14px; }
.wrap { overflow-x:auto; border:1px solid var(--line); border-radius:12px; }
table { width:100%; border-collapse:collapse; font-size:14.5px; }
th, td { text-align:left; vertical-align:top; padding:9px 10px; border-bottom:1px solid var(--line); }
th { position:sticky; top:0; background:var(--mint); cursor:pointer; white-space:nowrap; }
th[aria-sort="ascending"]::after { content:" \\25B2"; } th[aria-sort="descending"]::after { content:" \\25BC"; }
td.num { text-align:right; white-space:nowrap; }
details summary { cursor:pointer; color:var(--green); }
details ol { margin:6px 0 0 18px; padding:0; }
a { color:var(--green); }
.why { color:var(--soft); }
footer { color:var(--soft); font-size:13.5px; padding-top:18px; padding-bottom:28px; }
@media (max-width:640px) { #map { height:300px; } h1 { font-size:23px; } }
</style>
</head>
<body>
<header>
  <h1>Labs for: {{QUERY}}</h1>
  <p class="meta">{{YEARS}}. {{COUNTS}}. Created {{CREATED}}.</p>
</header>
<main>
  <div id="map" role="img" aria-label="Map of the labs that have a known place"></div>
  <div class="bar"><input id="q" type="search" placeholder="Filter by name, place, species or data" aria-label="Filter">
    <span id="shown"></span></div>
  <div class="wrap"><table id="labs">
    <thead><tr><th data-k="rank">#</th><th data-k="lead">Lead</th><th data-k="organization">Organisation</th>
      <th data-k="country">Country</th><th data-k="papers_as_senior_author">Senior</th><th data-k="last">Last</th>
      <th data-k="why">Why it came up</th></tr></thead>
    <tbody></tbody>
  </table></div>
</main>
<footer><p>{{NOTE}}</p><p>Map tiles: &copy; OpenStreetMap contributors. Made with Actogram Scholar {{VERSION}},
  <a href="https://github.com/actogram/scholar-mcp">github.com/actogram/scholar-mcp</a>.</p></footer>
<script type="application/json" id="data">{{DATA}}</script>
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"
  integrity="sha256-20nQCchB9co0qIjJZRGuk2/Z9VM+kNiyxNV1lvTlZBo=" crossorigin=""></script>
<script>
(function () {
  var data = JSON.parse(document.getElementById('data').textContent);
  var labs = data.labs.map(function (l, i) { l.rank = i + 1; l.last = (l.years || [])[1] || ''; return l; });
  function esc(s) { return String(s == null ? '' : s).replace(/[&<>"']/g, function (c) {
    return {'&': '&amp;', '<': '&lt;', '>': '&gt;', '"': '&quot;', "'": '&#39;'}[c]; }); }
  function papers(l) {
    if (!l.top_papers.length) return '';
    var n = l.top_papers.length;
    return '<details><summary>' + n + (n === 1 ? ' paper' : ' papers') + '</summary><ol>' + l.top_papers.map(function (p) {
      return '<li>' + (p.url ? '<a href="' + esc(p.url) + '" target="_blank" rel="noopener">' + esc(p.title) + '</a>' : esc(p.title)) +
        ' (' + esc(p.year) + ')</li>'; }).join('') + '</ol></details>';
  }
  var body = document.querySelector('#labs tbody');
  function row(l) {
    var lead = esc(l.lead) + (l.orcid ? ' <a href="https://orcid.org/' + esc(l.orcid) + '" target="_blank" rel="noopener">ORCID</a>' : '');
    var org = l.ror ? '<a href="' + esc(l.ror) + '" target="_blank" rel="noopener">' + esc(l.organization) + '</a>' : esc(l.organization);
    return '<tr><td class="num">' + l.rank + '</td><td>' + lead + '</td><td>' + org + (l.city ? '<br><span class="why">' +
      esc(l.city) + '</span>' : '') + '</td><td>' + esc(l.country) + '</td><td class="num">' + l.papers_as_senior_author +
      '</td><td class="num">' + esc(l.last) + '</td><td><span class="why">' + esc(l.why) + '</span>' + papers(l) + '</td></tr>';
  }
  var key = 'rank', dir = 1;
  function draw() {
    var q = document.getElementById('q').value.toLowerCase();
    var list = labs.filter(function (l) { return !q || [l.lead, l.organization, l.city, l.country, l.why].join(' ').toLowerCase().indexOf(q) >= 0; });
    list.sort(function (a, b) { var x = a[key], y = b[key]; return (x > y ? 1 : x < y ? -1 : 0) * dir; });
    body.innerHTML = list.map(row).join('');
    document.getElementById('shown').textContent = list.length + ' of ' + labs.length + ' labs';
  }
  document.querySelectorAll('#labs th').forEach(function (th) {
    th.addEventListener('click', function () {
      var k = th.getAttribute('data-k'); dir = (k === key) ? -dir : (k === 'rank' || k === 'lead' ? 1 : -1); key = k;
      document.querySelectorAll('#labs th').forEach(function (h) { h.removeAttribute('aria-sort'); });
      th.setAttribute('aria-sort', dir > 0 ? 'ascending' : 'descending'); draw();
    });
  });
  document.getElementById('q').addEventListener('input', draw);
  draw();
  if (window.L) {
    var map = L.map('map', { scrollWheelZoom: false }).setView([30, 10], 2);
    L.tileLayer('https://tile.openstreetmap.org/{z}/{x}/{y}.png', { maxZoom: 18,
      attribution: '&copy; <a href="https://www.openstreetmap.org/copyright">OpenStreetMap</a> contributors' }).addTo(map);
    var pts = [];
    labs.forEach(function (l) {
      if (l.lat == null || l.lon == null) return;
      var r = 5 + Math.min(10, 2 * l.papers_as_senior_author);
      L.circleMarker([l.lat, l.lon], { radius: r, color: '#00473a', weight: 1.5, fillColor: '#006956', fillOpacity: 0.55 })
        .bindPopup('<b>' + esc(l.lead) + '</b><br>' + esc(l.organization) + '<br>' + esc(l.why)).addTo(map);
      pts.push([l.lat, l.lon]);
    });
    if (pts.length > 1) map.fitBounds(pts, { padding: [24, 24], maxZoom: 6 });
    else if (pts.length === 1) map.setView(pts[0], 5);
  } else {
    document.getElementById('map').textContent = 'The map needs an internet connection; the table works without it.';
  }
})();
</script>
</body>
</html>
"""
